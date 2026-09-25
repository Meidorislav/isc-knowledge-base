import io

import docx
import pytest

from app.ingestion.parsing import UnsupportedFormatError, detect_format, normalize, parse


@pytest.mark.parametrize(
    ("filename", "expected"),
    [("a.txt", "txt"), ("Регламент.MD", "md"), ("x.markdown", "md"), ("b.docx", "docx")],
)
def test_detect_format(filename, expected):
    assert detect_format(filename) == expected


def test_detect_format_rejects_unknown():
    with pytest.raises(UnsupportedFormatError):
        detect_format("scan.pdf")


def test_parse_txt_cp1251():
    raw = "Порядок согласования договоров".encode("cp1251")
    assert parse(raw, "txt") == "Порядок согласования договоров"


def test_parse_txt_utf8_bom():
    raw = "﻿Привет".encode()
    assert parse(raw, "txt") == "Привет"


def test_normalize_collapses_blank_lines_and_crlf():
    assert normalize("a  \r\n\r\n\r\n\r\nb c\n") == "a\n\nb c"


def test_parse_docx_headings_and_tables():
    d = docx.Document()
    d.add_heading("Регламент закупок", level=1)
    d.add_paragraph("Общие положения.")
    d.add_heading("Сроки", level=2)
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Этап", "Срок"
    table.cell(1, 0).text, table.cell(1, 1).text = "Согласование", "3 дня"
    buf = io.BytesIO()
    d.save(buf)

    text = parse(buf.getvalue(), "docx")

    assert text.splitlines()[0] == "# Регламент закупок"
    assert "## Сроки" in text
    assert "| Этап | Срок |" in text
    assert "| Согласование | 3 дня |" in text
